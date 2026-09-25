"""Tab 7: inspect KiotViet export differences and resync saved locations."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from mapper.database import WarehouseDatabase
from mapper.common import application_directory
from mapper.location_comparison import (
    LocationDifference,
    compare_locations,
    export_comparison,
    read_web_locations,
    suggested_columns,
)
from mapper.web_upload import LocationUpdate


class WorkbookColumnDialog(tk.Toplevel):
    def __init__(self, parent: tk.Misc, workbook):
        super().__init__(parent)
        self.workbook = workbook
        self.result: tuple[str, int, int, int] | None = None
        self.title("Choose KiotViet Excel columns")
        self.resizable(False, False)
        self.transient(parent)
        self.grab_set()

        ttk.Label(self, text="Choose the worksheet, header row, and two columns from the export.").grid(
            row=0, column=0, columnspan=2, sticky="w", padx=14, pady=(14, 10)
        )
        ttk.Label(self, text="Worksheet").grid(row=1, column=0, sticky="w", padx=14, pady=5)
        self.sheet = ttk.Combobox(self, values=workbook.sheetnames, state="readonly", width=42)
        self.sheet.grid(row=1, column=1, padx=(0, 14), pady=5)
        self.sheet.current(0)
        self.sheet.bind("<<ComboboxSelected>>", self._update_columns)

        ttk.Label(self, text="Header row").grid(row=2, column=0, sticky="w", padx=14, pady=5)
        self.header_row = tk.IntVar(self, value=1)
        row_input = ttk.Spinbox(self, from_=1, to=100, width=7, textvariable=self.header_row)
        row_input.grid(row=2, column=1, sticky="w", pady=5)
        row_input.bind("<FocusOut>", self._update_columns)
        row_input.bind("<Return>", self._update_columns)

        ttk.Label(self, text="Product ID column").grid(row=3, column=0, sticky="w", padx=14, pady=5)
        self.product_column = ttk.Combobox(self, state="readonly", width=42)
        self.product_column.grid(row=3, column=1, padx=(0, 14), pady=5)
        ttk.Label(self, text="Website location column").grid(row=4, column=0, sticky="w", padx=14, pady=5)
        self.location_column = ttk.Combobox(self, state="readonly", width=42)
        self.location_column.grid(row=4, column=1, padx=(0, 14), pady=5)

        buttons = ttk.Frame(self)
        buttons.grid(row=5, column=0, columnspan=2, sticky="e", padx=14, pady=14)
        ttk.Button(buttons, text="Cancel", command=self.destroy).pack(side="right", padx=(8, 0))
        ttk.Button(buttons, text="Import", command=self._confirm).pack(side="right")
        self.protocol("WM_DELETE_WINDOW", self.destroy)
        self._suggest_header_row()
        self._update_columns()
        self.wait_visibility()
        self.focus_set()

    def _suggest_header_row(self) -> None:
        for sheet_name in self.workbook.sheetnames:
            worksheet = self.workbook[sheet_name]
            for row_number, row in enumerate(worksheet.iter_rows(max_row=15, values_only=True), 1):
                product, location = suggested_columns(list(row))
                if product != location and row[product] and row[location]:
                    from mapper.location_comparison import _plain_header, ID_HEADERS, LOCATION_HEADERS
                    if (_plain_header(row[product]) in ID_HEADERS
                            and _plain_header(row[location]) in LOCATION_HEADERS):
                        self.sheet.set(sheet_name)
                        self.header_row.set(row_number)
                        return

    def _update_columns(self, _event=None) -> None:
        try:
            row_number = self.header_row.get()
            if row_number < 1:
                return
            worksheet = self.workbook[self.sheet.get()]
            row = next(worksheet.iter_rows(min_row=row_number, max_row=row_number, values_only=True))
        except (ValueError, KeyError, StopIteration, tk.TclError):
            return
        labels = [f"{index + 1}: {str(value).strip() if value is not None else '(blank)'}"
                  for index, value in enumerate(row)]
        self.product_column.configure(values=labels)
        self.location_column.configure(values=labels)
        if labels:
            product, location = suggested_columns(list(row))
            self.product_column.current(product)
            self.location_column.current(location)

    def _confirm(self) -> None:
        try:
            header_row = self.header_row.get()
        except tk.TclError:
            header_row = 0
        product = self.product_column.current()
        location = self.location_column.current()
        if header_row < 1 or product < 0 or location < 0 or product == location:
            messagebox.showerror(
                "Choose columns", "Choose a valid header row and two different columns.", parent=self
            )
            return
        self.result = (self.sheet.get(), header_row, product, location)
        self.destroy()


class LocationComparisonView(ttk.Frame):
    def __init__(self, parent: tk.Misc, database: WarehouseDatabase, on_resync):
        super().__init__(parent, padding=10)
        self.database = database
        self.on_resync = on_resync
        self.web_locations: dict[str, str] | None = None
        self.imported_path: Path | None = None
        self.differences: list[LocationDifference] = []
        self.missing_ids: list[str] = []
        self.resynced: dict[str, str] = {}
        self.busy = False
        self.summary = tk.StringVar(self, value="Import a KiotViet Excel export to compare locations.")
        self.file_label = tk.StringVar(self, value="No web data imported")
        self.progress = tk.StringVar(self, value="")
        self._build_ui()

    def _build_ui(self) -> None:
        ttk.Label(self, text="Web Location Comparison", style="Heading.TLabel").pack(anchor="w")
        ttk.Label(
            self,
            text="Compare KiotViet locations with saved local shelf locations. Local locations are the true values.",
            foreground="#555555",
        ).pack(anchor="w", pady=(2, 10))
        actions = ttk.Frame(self)
        actions.pack(fill="x", pady=(0, 8))
        self.import_button = ttk.Button(actions, text="Import Web's Data", command=self.import_excel)
        self.import_button.pack(side="left")
        self.resync_button = ttk.Button(actions, text="Web resync", command=self.resync, state="disabled")
        self.resync_button.pack(side="left", padx=(8, 0))
        self.export_button = ttk.Button(
            actions, text="Export comparison excel", command=self.export_excel, state="disabled"
        )
        self.export_button.pack(side="left", padx=(8, 0))
        self.refresh_button = ttk.Button(actions, text="Refresh comparison", command=self.refresh)
        self.refresh_button.pack(side="left", padx=(8, 0))
        ttk.Label(self, textvariable=self.file_label, foreground="#555555").pack(anchor="w")
        ttk.Label(self, textvariable=self.summary).pack(anchor="w", pady=(4, 8))
        ttk.Label(self, textvariable=self.progress, foreground="#0969da").pack(anchor="w", pady=(0, 6))

        frame = ttk.Frame(self)
        frame.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(
            frame,
            columns=("product_id", "wrong_location", "true_location", "status"),
            show="headings",
        )
        for column, heading, width in (
            ("product_id", "Product ID", 190),
            ("wrong_location", "Website location", 260),
            ("true_location", "Local location", 260),
            ("status", "Status", 230),
        ):
            self.tree.heading(column, text=heading)
            self.tree.column(column, width=width, minwidth=100, anchor="w")
        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        scrollbar.pack(side="right", fill="y")
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.tag_configure("missing", foreground="#777777")
        self.tree.tag_configure("resynced", foreground="#217046")
        ttk.Label(
            self,
            text="Missing export IDs cannot be compared. Products with several or no local locations appear but need manual web review.",
            foreground="#555555",
        ).pack(anchor="w", pady=(8, 0))

    def _local_locations(self) -> dict[str, list[str]]:
        local: dict[str, list[str]] = {}
        registered_ids = {row[0] for row in self.database.get_product_records()}
        registered_ids.update(row[0] for row in self.database.get_catalog_products())
        for product_id, slot_name, _stock, _assigned_at in self.database.get_all_placement_items():
            if product_id not in registered_ids:
                continue
            slots = local.setdefault(product_id, [])
            if slot_name not in slots:
                slots.append(slot_name)
        return local

    def import_excel(self) -> None:
        if self.busy:
            return
        path = filedialog.askopenfilename(
            parent=self,
            title="Import KiotViet product locations",
            filetypes=(("Excel workbooks", "*.xlsx"),),
        )
        if not path:
            return
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                dialog = WorkbookColumnDialog(self, workbook)
                self.wait_window(dialog)
                if dialog.result is None:
                    return
                web_locations = read_web_locations(path, *dialog.result)
            finally:
                workbook.close()
        except Exception as error:
            messagebox.showerror("Excel import failed", str(error), parent=self)
            return
        self.web_locations = web_locations
        self.imported_path = Path(path)
        self.file_label.set(f"Imported {self.imported_path.name}: {len(web_locations):,} website product IDs")
        self.resynced.clear()
        self.progress.set("")
        self.refresh()

    def refresh(self) -> None:
        self.tree.delete(*self.tree.get_children())
        if self.web_locations is None:
            return
        local_locations = self._local_locations()
        self.differences, self.missing_ids = compare_locations(
            local_locations, self.web_locations
        )
        for difference in self.differences:
            if self.resynced.get(difference.product_id) == difference.true_location:
                status, tag = "Saved; re-export to verify", "resynced"
            elif not difference.can_resync:
                if not difference.true_location:
                    status, tag = "No local location; review", "missing"
                else:
                    status, tag = "Several local locations; review", "missing"
            else:
                status, tag = "Unsynced", ""
            self.tree.insert(
                "", "end", iid=f"difference::{difference.product_id}",
                values=(difference.product_id, difference.wrong_location,
                        difference.true_location, status),
                tags=(tag,) if tag else (),
            )
        for product_id in self.missing_ids:
            self.tree.insert(
                "", "end", iid=f"missing::{product_id}",
                values=(product_id, "(not in export)", ", ".join(local_locations[product_id]),
                        "Cannot compare"),
                tags=("missing",),
            )
        self._update_summary()

    def _update_summary(self) -> None:
        unsynced = self.exportable_differences()
        ready = sum(row.can_resync for row in unsynced)
        self.summary.set(
            f"{len(unsynced):,} unsynced · {ready:,} ready for web resync · "
            f"{len(self.missing_ids):,} local IDs absent from export"
        )
        self.resync_button.configure(state="normal" if ready and not self.busy else "disabled")
        self.export_button.configure(state="normal" if unsynced and not self.busy else "disabled")

    def exportable_differences(self) -> list[LocationDifference]:
        return [row for row in self.differences
                if self.resynced.get(row.product_id) != row.true_location]

    def pending_updates(self) -> list[LocationUpdate]:
        self.refresh()
        return [
            LocationUpdate(row.product_id, row.true_location)
            for row in self.differences
            if row.can_resync and self.resynced.get(row.product_id) != row.true_location
        ]

    def resync(self) -> None:
        updates = self.pending_updates()
        if not updates:
            messagebox.showinfo("No unsynced locations", "There are no comparable unsynced products to upload.", parent=self)
            return
        self.on_resync(updates)

    def mark_resynced(self, product_id: str) -> None:
        row = next((item for item in self.differences if item.product_id == product_id), None)
        if row is not None:
            self.resynced[product_id] = row.true_location
            iid = f"difference::{product_id}"
            if self.tree.exists(iid):
                self.tree.item(
                    iid,
                    values=(row.product_id, row.wrong_location, row.true_location,
                            "Saved; re-export to verify"),
                    tags=("resynced",),
                )
            self._update_summary()

    def set_busy(self, busy: bool) -> None:
        self.busy = busy
        self.import_button.configure(state="disabled" if busy else "normal")
        self.refresh_button.configure(state="disabled" if busy else "normal")
        self.refresh()

    def set_progress(self, message: str) -> None:
        self.progress.set(message)

    def export_excel(self) -> None:
        if self.busy:
            return
        self.refresh()
        differences = self.exportable_differences()
        if not differences:
            messagebox.showinfo("No differences", "There are no different locations to export.", parent=self)
            return
        try:
            result_dir = application_directory() / "xl" / "app_result"
            result_dir.mkdir(parents=True, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            path = result_dir / f"kiotviet_location_comparison_{timestamp}.xlsx"
            counter = 2
            while path.exists():
                path = result_dir / f"kiotviet_location_comparison_{timestamp}_{counter}.xlsx"
                counter += 1
            export_comparison(path, differences)
        except Exception as error:
            messagebox.showerror("Excel export failed", str(error), parent=self)
            return
        self.summary.set(f"Exported {len(differences):,} unsynced location(s) to {path}.")
