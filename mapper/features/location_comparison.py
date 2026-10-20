"""Read a KiotViet export and compare its locations with saved local placements."""

from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from mapper.core.common import ID_HEADERS, LOCATION_HEADERS, data_directory, normalize_sku, plain_header, split_locations


@dataclass(frozen=True)
class LocationDifference:
    product_id: str
    wrong_location: str
    true_location: str


def is_header_row(row: list[object]) -> bool:
    """True when the row names both a product ID column and a location column."""
    keys = [plain_header(value) for value in row]
    return any(key in ID_HEADERS for key in keys) and any(key in LOCATION_HEADERS for key in keys)


def export_path(directory: Path | None = None, now: datetime | None = None) -> Path:
    """A new timestamped file in data/exports/location_comparison."""
    directory = directory or data_directory() / "exports" / "location_comparison"
    directory.mkdir(parents=True, exist_ok=True)
    stem = f"kiotviet_location_comparison_{(now or datetime.now()).strftime('%Y%m%d_%H%M%S')}"
    path, counter = directory / f"{stem}.xlsx", 2
    while path.exists():
        path, counter = directory / f"{stem}_{counter}.xlsx", counter + 1
    return path


def read_web_locations(
    path: Path | str,
    sheet_name: str,
    header_row: int,
    id_column: int,
    location_column: int,
) -> dict[str, str]:
    """Import the chosen worksheet without changing the local database."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Excel import needs openpyxl in the Python environment."
        ) from error

    if (
        header_row < 1
        or id_column < 0
        or location_column < 0
        or id_column == location_column
    ):
        raise ValueError(
            "Choose different product ID and location columns and a valid header row."
        )
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet_name!r} is not in the Excel file.")
        worksheet = workbook[sheet_name]
        locations: dict[str, str] = {}
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            header_row + 1,
        ):
            product = row[id_column] if id_column < len(row) else None
            if product is None or not str(product).strip():
                continue
            if isinstance(product, float) and product.is_integer():
                product = int(product)
            product_id = normalize_sku(product)
            value = row[location_column] if location_column < len(row) else None
            location = str(value).strip() if value is not None else ""
            if product_id in locations:
                # A duplicate row with a blank location adds nothing; keep the located one.
                if not location:
                    continue
                previous = locations[product_id]
                if previous and set(split_locations(previous)) != set(split_locations(location)):
                    raise ValueError(
                        f"Product {product_id} has conflicting website locations at Excel row {row_number}."
                    )
            locations[product_id] = location
        if not locations:
            raise ValueError(
                "No product IDs were found in the selected worksheet and column."
            )
        return locations
    finally:
        workbook.close()


def compare_locations(
    local_locations: dict[str, Iterable[str]], web_locations: dict[str, str]
) -> tuple[list[LocationDifference], list[str]]:
    """Return mismatches and local assigned IDs missing from the export; keys are canonical SKUs."""
    differences: list[LocationDifference] = []
    missing: list[str] = []
    for product_id in sorted(local_locations.keys() | web_locations.keys(), key=str.casefold):
        local = split_locations(local_locations.get(product_id, ()))
        if product_id not in web_locations:
            if local:
                missing.append(product_id)
        elif set(local) != set(split_locations(web_locations[product_id])):
            differences.append(LocationDifference(product_id, web_locations[product_id], ", ".join(local)))
    return differences, missing


def export_comparison(path: Path | str, differences: list[LocationDifference]) -> None:
    """Write only the three requested columns, keeping all IDs as text."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as error:
        raise RuntimeError(
            "Excel export needs openpyxl in the Python environment."
        ) from error

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Unsynced locations"
    worksheet.append(("product_id", "wrong_location", "true_location"))
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24476B")
    for difference in differences:
        worksheet.append(
            (difference.product_id, difference.wrong_location, difference.true_location)
        )
        for cell in worksheet[worksheet.max_row]:
            cell.data_type = "s"  # an ID such as "=SKU" must stay text, never a formula
    worksheet.column_dimensions["A"].width = 25
    worksheet.column_dimensions["B"].width = 36
    worksheet.column_dimensions["C"].width = 36
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    workbook.save(path)
