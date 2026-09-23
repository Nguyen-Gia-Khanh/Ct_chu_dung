"""Read a KiotViet export and compare its locations with saved local placements."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
import unicodedata


@dataclass(frozen=True)
class LocationDifference:
    product_id: str
    wrong_location: str
    true_location: str
    can_resync: bool


def _plain_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").casefold())
    return re.sub(r"[^a-z0-9]", "", "".join(c for c in text if not unicodedata.combining(c)))


ID_HEADERS = (
    "mahang", "masanpham", "masp", "productid", "productcode", "sku", "itemcode", "code"
)
LOCATION_HEADERS = (
    "vitri", "vitrikho", "vitriluutru", "vitriluutrukho", "locationid", "location",
    "warehouselocation", "shelflocation", "shelf"
)


def suggested_columns(headers: list[object]) -> tuple[int, int]:
    normalized = [_plain_header(header) for header in headers]
    product_index = next((normalized.index(name) for name in ID_HEADERS if name in normalized), 0)
    location_index = next(
        (normalized.index(name) for name in LOCATION_HEADERS if name in normalized),
        1 if len(headers) > 1 else 0,
    )
    return product_index, location_index


def location_parts(value: object) -> tuple[str, ...]:
    """KiotViet can export several shelf IDs in one cell."""
    if value is None:
        return ()
    text = str(value).strip()
    return tuple(dict.fromkeys(part.strip() for part in re.split(r"[,;\n]+", text) if part.strip()))


def read_web_locations(
    path: Path | str, sheet_name: str, header_row: int, id_column: int, location_column: int
) -> dict[str, str]:
    """Import the chosen worksheet without changing the local database."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Excel import needs openpyxl in the Python environment.") from error

    if header_row < 1 or id_column < 0 or location_column < 0 or id_column == location_column:
        raise ValueError("Choose different product ID and location columns and a valid header row.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet_name!r} is not in the Excel file.")
        worksheet = workbook[sheet_name]
        locations: dict[str, str] = {}
        for row_number, row in enumerate(worksheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            product = row[id_column] if id_column < len(row) else None
            if product is None or not str(product).strip():
                continue
            if isinstance(product, float) and product.is_integer():
                product = int(product)
            product_id = str(product).strip()
            value = row[location_column] if location_column < len(row) else None
            location = str(value).strip() if value is not None else ""
            if product_id in locations and set(location_parts(locations[product_id])) != set(location_parts(location)):
                raise ValueError(
                    f"Product {product_id} has conflicting website locations at Excel row {row_number}."
                )
            locations[product_id] = location
        if not locations:
            raise ValueError("No product IDs were found in the selected worksheet and column.")
        return locations
    finally:
        workbook.close()


def compare_locations(
    local_locations: dict[str, list[str]], web_locations: dict[str, str]
) -> tuple[list[LocationDifference], list[str]]:
    """Return mismatches and local assigned IDs missing from the export."""
    differences: list[LocationDifference] = []
    missing: list[str] = []
    for product_id in sorted(local_locations, key=str.casefold):
        local = location_parts(", ".join(local_locations[product_id]))
        if not local:
            continue
        if product_id not in web_locations:
            missing.append(product_id)
            continue
        web = location_parts(web_locations[product_id])
        if set(local) != set(web):
            differences.append(
                LocationDifference(
                    product_id,
                    web_locations[product_id],
                    ", ".join(local),
                    len(local) == 1,
                )
            )
    return differences, missing


def export_comparison(path: Path | str, differences: list[LocationDifference]) -> None:
    """Write only the three requested columns, keeping all IDs as text."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as error:
        raise RuntimeError("Excel export needs openpyxl in the Python environment.") from error

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Unsynced locations"
    worksheet.append(("product_id", "wrong_location", "true_location"))
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24476B")
    for difference in differences:
        worksheet.append((difference.product_id, difference.wrong_location, difference.true_location))
        for cell in worksheet[worksheet.max_row]:
            cell.data_type = "s"
    worksheet.column_dimensions["A"].width = 25
    worksheet.column_dimensions["B"].width = 36
    worksheet.column_dimensions["C"].width = 36
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    workbook.save(path)
