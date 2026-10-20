"""KiotViet Excel comparison and selective resync checks."""

from pathlib import Path
from queue import Empty, Queue
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import Mock

from openpyxl import Workbook, load_workbook

from mapper.features.location_comparison import (
    LocationDifference,
    compare_locations,
    export_comparison,
    read_web_locations,
)
from ui.desktop.app import WarehouseMapperApp
from ui.desktop.location_comparison_view import LocationComparisonView
from mapper.features.web_upload import LocationUpdate


class LocationComparisonTests(unittest.TestCase):
    def test_kiotviet_headers_and_excel_import(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "kiotviet.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(("Loại hàng", "Nhóm hàng(3 Cấp)", "Mã hàng", "Mã vạch", "Tên hàng", "Vị trí"))
            sheet.append(("Item", "Group", "SKU-1", "barcode", "Name", "L1-1A1-1"))
            sheet.append(("Item", "Group", "SKU-2", "barcode", "Name", None))
            sheet.append(("Item", "Group", "WEB-ONLY", "barcode", "Name", "L9-9Z9-9"))
            workbook.save(path)

            self.assertEqual(
                read_web_locations(path, sheet.title, 1, 2, 5),
                {"SKU-1": "L1-1A1-1", "SKU-2": "", "WEB-ONLY": "L9-9Z9-9"},
            )

    def test_duplicate_ids_ignore_blank_rows_but_reject_conflicts(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "kiotviet.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(("Mã hàng", "Vị trí"))
            sheet.append(("BLANKFIRST", None))
            sheet.append(("BLANKFIRST", "L1-1A1-1"))
            sheet.append(("BLANKLAST", "L1-1A1-2"))
            sheet.append(("BLANKLAST", None))
            sheet.append(("BOTHBLANK", None))
            sheet.append(("BOTHBLANK", None))
            workbook.save(path)

            self.assertEqual(
                read_web_locations(path, sheet.title, 1, 0, 1),
                {"BLANKFIRST": "L1-1A1-1", "BLANKLAST": "L1-1A1-2", "BOTHBLANK": ""},
            )

            sheet.append(("BLANKFIRST", "L2-2B2-2"))
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "BLANKFIRST has conflicting website locations at Excel row 8"):
                read_web_locations(path, sheet.title, 1, 0, 1)

    def test_comparison_checks_both_ways_and_identifies_missing(self):
        local = {
            "SKU-1": ["L1-1A1-1"],
            "SKU-2": ["L1-1A1-2"],
            "SKU-3": ["L1-1A1-3"],
            "MULTI": ["L1-1A1-1", "L1-1A1-2"],
        }
        web = {
            "SKU-1": " L1-1A1-1 ",
            "SKU-2": "",
            "MULTI": "L1-1A1-1",
            "WEB-ONLY": "L9-9Z9-9",
        }
        differences, missing = compare_locations(local, web)
        self.assertEqual(
            [item.product_id for item in differences], ["MULTI", "SKU-2", "WEB-ONLY"]
        )
        self.assertEqual(differences[2].wrong_location, "L9-9Z9-9")
        self.assertEqual(differences[2].true_location, "")
        self.assertEqual(missing, ["SKU-3"])

    def test_numeric_export_ids_regain_the_saved_leading_zeros(self):
        local = {
            "06451961405": ["L1-3D2-14"],
            "0123": ["L1-1A1-1"],
            "00123": ["L1-1A1-2"],
            "0777": ["L1-1A1-3"],
            "777": ["L1-1A1-4"],
        }
        web = {"6451961405": "L1-3D2-14", "123": "L1-1A1-1", "777": "L1-1A1-4"}
        differences, missing = compare_locations(local, web)
        # 06451961405 matches; 123 is ambiguous; 777 already matches exactly.
        self.assertEqual([item.product_id for item in differences], ["123"])
        self.assertEqual(missing, ["00123", "0123", "0777"])

    def test_web_occupied_but_locally_unassigned_is_unsynced(self):
        local = {}
        web = {"PROD-100": "A-01-01", "PROD-200": ""}
        differences, missing = compare_locations(local, web)
        self.assertEqual(len(differences), 1)
        self.assertEqual(differences[0].product_id, "PROD-100")
        self.assertEqual(differences[0].wrong_location, "A-01-01")
        self.assertEqual(differences[0].true_location, "")
        self.assertEqual(missing, [])

    def test_tab_queues_web_location_clear_for_unassigned_product(self):
        differences, _ = compare_locations({}, {"PROD-100": "A-01-01"})
        view = SimpleNamespace(refresh=Mock(), differences=differences, resynced={})

        self.assertEqual(
            LocationComparisonView.pending_updates(view),
            [LocationUpdate("PROD-100", "")],
        )

    def test_tab_marks_cleared_location_as_saved(self):
        differences, _ = compare_locations({}, {"PROD-100": "A-01-01"})
        tree = Mock()
        tree.exists.return_value = True
        view = SimpleNamespace(
            differences=differences, resynced={}, tree=tree, _update_summary=Mock()
        )

        LocationComparisonView.mark_resynced(view, "PROD-100")

        self.assertEqual(view.resynced, {"PROD-100": ""})
        self.assertEqual(tree.item.call_args.kwargs["values"][2], "(none)")

    def test_comparison_worker_queues_data_without_touching_widgets(self):
        events = Queue()
        view = SimpleNamespace(
            _comparison_events=events,
            database=Mock(get_location_sets=Mock(return_value={"P1": ("A-02",)})),
            tree=Mock(),
        )

        LocationComparisonView._run_comparison(
            view, "refresh", None, None, {"P1": "A-01"}, {}
        )

        event = events.get_nowait()
        self.assertEqual(event[0], "done")
        self.assertEqual(event[3][0].true_location, "A-02")
        view.tree.insert.assert_not_called()

    def test_comparison_rows_render_in_small_ui_batches(self):
        callbacks = []
        differences = [LocationDifference(f"P{i}", "A-01", "A-02") for i in range(450)]
        view = SimpleNamespace(
            _render_rows=[("difference", row) for row in differences],
            _render_index=0,
            _render_local_locations={},
            _comparison_busy=True,
            _comparison_action="refresh",
            resynced={},
            tree=Mock(),
            after=Mock(side_effect=lambda _delay, callback: callbacks.append(callback)),
            _set_busy_state=Mock(),
            progress=Mock(),
        )
        view._render_comparison_rows = lambda: LocationComparisonView._render_comparison_rows(view)

        view._render_comparison_rows()
        self.assertEqual(view.tree.insert.call_count, 200)
        while callbacks:
            callbacks.pop(0)()

        self.assertEqual(view.tree.insert.call_count, 450)
        self.assertFalse(view._comparison_busy)
        self.assertEqual(view.after.call_count, 2)

    def test_export_has_only_requested_columns_and_preserves_text(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "comparison.xlsx"
            export_comparison(
                path,
                [LocationDifference("=SKU-1", "L0-1A1-1", "L1-1A1-1")],
            )
            workbook = load_workbook(path)
            sheet = workbook.active
            self.assertEqual(sheet.max_column, 3)
            self.assertEqual(
                list(sheet.values),
                [("product_id", "wrong_location", "true_location"),
                 ("=SKU-1", "L0-1A1-1", "L1-1A1-1")],
            )
            self.assertEqual(sheet["A2"].data_type, "s")

    def test_tab_resync_uses_location_only_batch_with_forced_saves(self):
        app = WarehouseMapperApp.__new__(WarehouseMapperApp)
        app._can_start_web_update = Mock(return_value=True)
        app._start_web_update = Mock()
        updates = [LocationUpdate("SKU-2", "L1-1A1-2")]

        app.resync_web_locations(updates)

        app._start_web_update.assert_called_once_with(
            updates, location_only=True, force_save_each_product=True, source="comparison"
        )

    def test_failed_batch_reports_only_completed_product_ids(self):
        app = WarehouseMapperApp.__new__(WarehouseMapperApp)
        app.website_uploader = Mock()
        app.website_uploader.modify_location.side_effect = ["saved", RuntimeError("failed")]
        app.website_uploader.is_connected.return_value = True
        events = Queue()
        updates = [LocationUpdate("SKU-1", "L1-1A1-1"),
                   LocationUpdate("SKU-2", "L1-1A1-2")]

        app._run_web_upload(updates, events, location_only=True, save_each_product=True)

        results = []
        while True:
            try:
                results.append(events.get_nowait())
            except Empty:
                break
        self.assertEqual([entry[1] for entry in results if entry[0] == "item_done"], ["SKU-1"])
        self.assertEqual(results[-1][0], "error")
        self.assertEqual(app.website_uploader.modify_location.call_count, 2)

if __name__ == "__main__":
    unittest.main()
