"""Check automatic photo layout, wrapped rows, and text-safe Excel output."""

from pathlib import Path
import runpy
import sys
import unittest
from unittest.mock import patch

import cv2
import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from tests.support import WorkspaceTestCase
from utils import statement_ocr
from utils.statement_ocr import COLUMNS, assemble, detect_rows, detect_table, discover_images, extract, write_excel


class StatementOCRTests(WorkspaceTestCase):
    def test_manual_run_uses_local_ocr_environment_without_options(self):
        script = Path(__file__).resolve().parents[1] / "utils" / "statement_ocr.py"
        interpreter = script.parent.parent / ".venv-ocr" / "Scripts" / "python.exe"
        with patch.object(sys, "executable", "other-python.exe"), patch.object(Path, "is_file", return_value=True):
            with patch("subprocess.call", return_value=0) as launch:
                with self.assertRaises(SystemExit) as stopped:
                    runpy.run_path(str(script), run_name="__main__")
                self.assertEqual(stopped.exception.code, 0)
                launch.assert_called_once_with([str(interpreter), str(script)])

    def test_each_photo_detects_its_own_perspective_and_columns(self):
        image = Image.new("RGB", (900, 800), "white")
        draw = ImageDraw.Draw(image)
        dividers = [30, 70, 145, 220, 400, 640, 715, 790, 870]
        for x in dividers:
            # Broken rules like the photographed statements, without row lines.
            for y in range(60, 420 if x in (220, 400, 640) else 740, 80):
                draw.line((x, y, x, y + 64), fill="black", width=3)
        source = np.float32([[0, 0], [899, 0], [899, 799], [0, 799]])
        for corners in (source, np.float32([[60, 80], [960, 40], [930, 870], [25, 800]])):
            matrix = cv2.getPerspectiveTransform(source, corners)
            photo = Image.fromarray(cv2.warpPerspective(np.asarray(image), matrix, (1000, 920), borderValue=(255, 255, 255)))
            table, layout = detect_table(photo)
            self.assertEqual(len(layout["columns"]), 9)
            proportions = np.array(layout["columns"]) / table.width
            np.testing.assert_allclose(proportions, (np.array(dividers) - 30) / 840, atol=0.015)
        with self.assertRaisesRegex(ValueError, "column boundaries"):
            detect_table(Image.new("RGB", (900, 800), "white"))

    def test_wrapped_rows_missing_index_and_excel_text_are_preserved(self):
        def token(text, top, bottom=None, score=0.99):
            return {"text": text, "score": score, "box": [2, top, 25, bottom or top + 12]}

        columns = [[token("001", 10), token("003", 135)],
                   [token("17/06/2025 09:17:14", 10, 38), token("17/06/2025", 95), token("18/06/2025", 135)],
                   [token("17/06/2025", 11), token("17/06/2025", 96), token("18/06/2025", 136)],
                   [token("000123", 10)],
                   [token("first", 10), token("wrapped", 65), token("=formula", 95)],
                   [], [token("1,000,000", 10, score=0.65)], [token("9,001,042", 10)]]
        rows = detect_rows(columns, 180)
        result = assemble(columns, rows, 0.9)
        self.assertEqual(len(result["rows"]), 3)
        self.assertEqual(result["rows"][0]["content"], "first wrapped")
        self.assertEqual(result["rows"][1]["index"], "")
        self.assertEqual(result["rows"][1]["content"], "=formula")
        self.assertTrue(any(issue["issue"] == "low_confidence" for issue in result["issues"]))
        predictions = []
        for column, tokens in enumerate(columns):
            scale = 3 if column == 0 else 1
            predictions.append({"rec_texts": [t["text"] for t in tokens],
                                "rec_scores": [t["score"] for t in tokens],
                                "rec_boxes": [[12, t["box"][1] * scale + 12,
                                               35, t["box"][3] * scale + 12] for t in tokens]})

        class FakeOCR:
            def predict(self, crop):
                return [predictions.pop(0)]

        layout = {"columns": [0, 40, 120, 200, 340, 550, 640, 730, 820]}
        extracted = extract(Image.new("RGB", (820, 180), "white"), layout, 0.9, FakeOCR())
        self.assertEqual(extracted["rows"], result["rows"])
        with self.assertRaisesRegex(ValueError, "No transaction"):
            detect_rows([[] for _ in COLUMNS], 180)
        folder = self.folder
        for name in ("page10.png", "page2.JPG", "readme.txt"):
            (folder / name).touch()
        self.assertEqual([path.name for path in discover_images(folder)], ["page2.JPG", "page10.png"])
        result["source"] = "page2.JPG"
        output = folder / "result.xlsx"
        write_excel(output, [result], [{"source_file": "bad.jpg", "error": "Unreadable image"}])
        workbook = load_workbook(output)
        try:
            self.assertEqual(workbook.sheetnames, ["Data"])
            self.assertEqual(workbook["Data"].max_column, 9)
            self.assertEqual(workbook["Data"]["I1"].value, "flagged")
            self.assertEqual(workbook["Data"]["D2"].value, "000123")
            self.assertEqual(workbook["Data"]["E3"].value, "=formula")
            self.assertEqual(workbook["Data"]["E3"].data_type, "s")
            self.assertIsNone(workbook["Data"]["F2"].value)
            self.assertIn("money_in: low confidence 0.65 < 0.90", workbook["Data"]["I2"].value)
            self.assertIn("missing or non numeric index", workbook["Data"]["I3"].value)
            self.assertEqual(workbook["Data"]["I5"].value, "bad.jpg: Unreadable image")
            self.assertEqual(workbook["Data"].auto_filter.ref, "A1:I5")
        finally:
            workbook.close()
        with self.assertRaises(FileExistsError):
            write_excel(output, [result], [])

    def test_confidence_cutoff_ignores_small_weak_fragments(self):
        columns = [[{"text": text, "score": 0.99, "box": [2, 10, 25, 22]}] for text in
                   ("001", "09:17:14", "17/06/2025", "000123", "long readable description", "", "1,000", "9,000")]
        columns[5] = []  # A blank outgoing amount is valid when incoming has a value.
        columns[4][0]["score"] = 0.95
        columns[4].append({"text": ".", "score": 0.2, "box": [26, 10, 30, 22]})
        for score in (0.85, 0.80, 0.79):
            columns[6][0]["score"] = score
            result = assemble(columns, [0, 40], 0.80)
            self.assertGreater(result["confidence"][0]["content"], 0.90)
            self.assertEqual([issue["column"] for issue in result["issues"]], ["money_in"] if score < 0.8 else [])
        columns[6][0]["score"] = 0.99
        columns[3] = []
        self.assertEqual(assemble(columns, [0, 40], 0.8)["issues"],
                         [{"row": 1, "column": "id", "issue": "no_text_detected"}])

    def test_batch_writes_only_one_workbook_and_keeps_failures_visible(self):
        folder = self.folder
        images = folder / "images"
        images.mkdir()
        for name in ("page1.png", "page2.png"):
            (images / name).touch()
        result = {"rows": [dict.fromkeys(COLUMNS, "001")], "issues": []}
        with patch.multiple(statement_ocr, SCRIPT_DIR=folder, IMAGE_FOLDER="images", OUTPUT_FILENAME="result.xlsx"):
            with patch.object(statement_ocr, "create_ocr"), patch.object(statement_ocr, "prepare_table", return_value=(None, {})):
                with patch.object(statement_ocr, "extract", side_effect=[ValueError("Unreadable image"), result]):
                    statement_ocr.main()
        self.assertEqual(sorted(path.name for path in folder.iterdir()), ["images", "result.xlsx"])
        workbook = load_workbook(folder / "result.xlsx")
        try:
            self.assertEqual(workbook.sheetnames, ["Data"])
            self.assertIsNone(workbook["Data"]["I2"].value)
            self.assertEqual(workbook["Data"]["I3"].value, "page1.png: Failed: Unreadable image")
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
