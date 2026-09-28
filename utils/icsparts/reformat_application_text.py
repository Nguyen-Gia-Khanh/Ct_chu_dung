"""Refresh part_master.application_text from the already-correct application_dict."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = (
    Path(__file__).resolve().parents[2]
    / "data"
    / "icsparts"
    / "icsparts_catalog_data.xlsx"
)


def format_application_text(value: str | None) -> str:
    if not value:
        return ""
    applications = json.loads(value)
    return "\n".join(
        f"{archetype}: "
        + "; ".join(
            f"{model} [{', '.join(versions)}]"
            for model, versions in models.items()
        )
        for archetype, models in applications.items()
    )


def main() -> None:
    workbook = load_workbook(WORKBOOK)
    sheet = workbook["part_master"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    dict_column = columns["application_dict"]
    text_column = columns["application_text"]

    updated = 0
    for row in range(2, sheet.max_row + 1):
        value = sheet.cell(row, dict_column).value
        text = format_application_text(value)
        if sheet.cell(row, text_column).value != text:
            sheet.cell(row, text_column).value = text
            updated += 1

    temporary = WORKBOOK.with_name(f".{WORKBOOK.stem}.application_text.tmp.xlsx")
    try:
        workbook.save(temporary)
        temporary.replace(WORKBOOK)
    finally:
        if temporary.exists():
            temporary.unlink()

    print(f"Updated {updated} application_text cells in {WORKBOOK}")


if __name__ == "__main__":
    main()
