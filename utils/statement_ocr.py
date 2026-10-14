"""Read images beside this script in images/ and export an Excel file with a filterable flagged column."""

from bisect import bisect_right
from datetime import datetime
import os
from pathlib import Path
import re
import statistics
import sys
from uuid import uuid4

COLUMNS = ("index", "time", "date", "id", "content", "money_out", "money_in", "total_raw")
SCRIPT_DIR = Path(__file__).resolve().parent

IMAGE_FOLDER = "images"
OUTPUT_FILENAME = ""  # Empty: timestamped workbook beside this script.
MIN_CONFIDENCE = 0.80  # Flag cells below this OCR score (0-1).

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}


def discover_images(folder):
    def natural_key(path):
        return tuple((1, int(part)) if part.isdigit() else (0, part.casefold())
                     for part in re.split(r"(\d+)", path.name))
    return sorted((p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in IMAGE_SUFFIXES), key=natural_key)


def validate_edges(edges, extent, expected=None):
    if expected is not None and len(edges) != expected:
        raise ValueError(f"Expected {expected} boundaries, got {len(edges)}")
    if len(edges) < 2 or edges[0] != 0 or edges[-1] != extent:
        raise ValueError("Boundaries must span the entire rectified table")
    if any(b - a < 3 for a, b in zip(edges, edges[1:])):
        raise ValueError("Boundaries must be increasing and at least three pixels apart")


def assemble(column_results, row_edges, min_confidence):
    """Group OCR boxes into detected rows without inventing missing values."""
    cells = [[[] for _ in COLUMNS] for _ in range(len(row_edges) - 1)]
    issues = []
    for column, tokens in enumerate(column_results):
        for token in tokens:
            x0, y0, x1, y1 = token["box"]
            row = bisect_right(row_edges, (y0 + y1) / 2) - 1
            if not 0 <= row < len(cells):
                issues.append({"column": COLUMNS[column], "issue": "outside_rows", "token": token})
                continue
            cells[row][column].append(token)
            # Small overlap from detector padding is tolerated; larger overlap needs review.
            margin = max(2, (y1 - y0) * 0.15)
            if y0 < row_edges[row] - margin or y1 > row_edges[row + 1] + margin:
                issues.append({"row": row + 1, "column": COLUMNS[column], "issue": "box_crosses_row", "token": token})
    records, confidence = [], []
    for row_number, row in enumerate(cells, 1):
        record, scores = {}, {}
        for name, tokens in zip(COLUMNS, row):
            # ponytail: near-horizontal text only; clustered baselines for more complex layouts.
            tokens.sort(key=lambda t: (round(t["box"][1] / 5), t["box"][0]))
            record[name] = " ".join(t["text"] for t in tokens).strip()
            # Weight each text line by its length, so tiny uncertain fragments do not dominate.
            characters = sum(len(t["text"].strip()) for t in tokens)
            scores[name] = sum(t["score"] * len(t["text"].strip()) for t in tokens) / characters if characters else None
            if scores[name] is not None and scores[name] < min_confidence:
                issues.append({"row": row_number, "column": name, "issue": "low_confidence",
                               "score": scores[name], "threshold": min_confidence})
        if not record["index"].isdigit():
            issues.append({"row": row_number, "column": "index", "issue": "missing_or_non_numeric_index"})
        for name in ("time", "date", "id", "content", "total_raw"):
            if not record[name]:
                issues.append({"row": row_number, "column": name, "issue": "no_text_detected"})
        if not record["money_out"] and not record["money_in"]:
            issues.append({"row": row_number, "issue": "no_amount_detected"})
        if record["money_out"] and record["money_in"]:
            issues.append({"row": row_number, "issue": "both_amount_columns_populated"})
        records.append(record)
        confidence.append(scores)
    indices = [int(r["index"]) if r["index"].isdigit() else None for r in records]
    for row, (previous, current) in enumerate(zip(indices, indices[1:]), 2):
        if previous is not None and current is not None and current != previous + 1:
            issues.append({"row": row, "issue": "index_gap_or_duplicate", "previous": previous, "current": current})
    return {"rows": records, "confidence": confidence, "issues": issues}


def write_excel(path, results, failures=()):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    if path.exists():
        raise FileExistsError(f"Output exists: {path}")
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append((*COLUMNS, "flagged"))
    for result in results:
        flags = {}
        for issue in result["issues"]:
            message = issue["issue"].replace("_", " ")
            if issue.get("column"):
                message = issue["column"] + ": " + message
            if issue["issue"] == "low_confidence":
                message += f" {issue['score']:.2f} < {issue['threshold']:.2f}"
            if "row" not in issue:
                message += ": " + issue.get("token", {}).get("text", "")
            flags.setdefault(issue.get("row"), []).append(message)
        for number, record in enumerate(result["rows"], 1):
            messages = flags.get(number, [])
            flagged = result["source"] + ": " + "; ".join(dict.fromkeys(messages)) if messages else None
            data.append((*[record[name] or None for name in COLUMNS], flagged))
            data.row_dimensions[data.max_row].height = min(160, max(30, 15 * (len(record["content"]) // 75 + 1)))
        if flags.get(None):
            data.append((*[None for _ in COLUMNS], result["source"] + ": " + "; ".join(dict.fromkeys(flags[None]))))
    # Keep failed photos visible when filtering flagged, without creating another sheet.
    for failure in failures:
        data.append((*[None for _ in COLUMNS], failure["source_file"] + ": " + failure["error"]))
    data.freeze_panes = "A2"
    data.auto_filter.ref = data.dimensions
    data.row_dimensions[1].height = 24
    for cell in data[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for column, width in enumerate((9, 24, 16, 24, 78, 20, 20, 20, 100), 1):
        data.column_dimensions[data.cell(1, column).column_letter].width = width
    for row in data.iter_rows(min_row=2):
        for cell in row:
            cell.data_type = "s"  # Preserve IDs and prevent OCR strings becoming formulas.
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row[-1].value:
            row[-1].fill = PatternFill("solid", fgColor="FFF2CC")
    temporary = path.with_name(path.stem + ".tmp_" + uuid4().hex + ".xlsx")
    try:
        workbook.save(temporary)
        if path.exists():
            raise FileExistsError(f"Output exists: {path}")
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def rectify(image, corners):
    import cv2
    import numpy as np
    from PIL import Image

    points = np.asarray(corners, dtype=np.float32)
    if points.shape != (4, 2) or not cv2.isContourConvex(points.astype(np.int32)):
        raise ValueError("Could not determine a convex table outline")
    width = round(max(np.linalg.norm(points[1] - points[0]), np.linalg.norm(points[2] - points[3])))
    height = round(max(np.linalg.norm(points[3] - points[0]), np.linalg.norm(points[2] - points[1])))
    if min(width, height) < 30:
        raise ValueError("Detected table is too small")
    target = np.float32([[0, 0], [width - 1, 0], [width - 1, height - 1], [0, height - 1]])
    matrix = cv2.getPerspectiveTransform(points, target)
    return Image.fromarray(cv2.warpPerspective(np.asarray(image), matrix, (width, height)))


def detect_table(image):
    """Find this photo's nine vertical rules and correct its perspective."""
    import cv2
    import numpy as np

    scale = min(1, 1200 / max(image.size))
    photo = cv2.resize(np.asarray(image), (round(image.width * scale), round(image.height * scale)))
    height, width = photo.shape[:2]
    gray = cv2.cvtColor(photo, cv2.COLOR_RGB2GRAY)
    binary = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                   cv2.THRESH_BINARY_INV, 31, 10)
    # Ignore colored screen borders; keep the dark printed rules.
    binary[cv2.cvtColor(photo, cv2.COLOR_RGB2HSV)[:, :, 1] > 80] = 0
    binary = cv2.morphologyEx(binary, cv2.MORPH_OPEN,
                             np.ones((max(5, height // 80), 1), np.uint8))
    lines = cv2.HoughLinesP(binary, 1, np.pi / 720, threshold=max(20, height // 27),
                           minLineLength=height // 10, maxLineGap=height // 30)
    candidates = []
    for x0, y0, x1, y1 in (() if lines is None else lines.reshape(-1, 4)):
        if abs(y1 - y0) > height * 0.15 and abs(x1 - x0) < abs(y1 - y0) * 0.25:
            middle = x0 + (height / 2 - y0) * (x1 - x0) / (y1 - y0)
            candidates.append((middle, (x0, y0, x1, y1)))
    groups = []
    for middle, line in sorted(candidates):
        if not groups or middle - statistics.mean(item[0] for item in groups[-1]) > max(3, width * 0.008):
            groups.append([])
        groups[-1].append((middle, line))
    if len(groups) != 9:
        raise ValueError(f"Found {len(groups)} column boundaries; expected 9 for this eight-column statement. "
                         "Use a clearer photo showing all the vertical dividers.")
    rules, tops, bottoms = [], [], []
    for group in groups:
        points = np.array([(x, y) for _, (x0, y0, x1, y1) in group
                           for x, y in ((x0, y0), (x1, y1))], dtype=float)
        rules.append(np.polyfit(points[:, 1], points[:, 0], 1))
        tops.append(float(points[:, 1].min()))
        bottoms.append(float(points[:, 1].max()))
    middle_x = np.array([np.polyval(rule, height / 2) for rule in rules])

    def horizontal_edge(ends, padding):
        # Broken rules can end early. Fit the boundary supported by most dividers.
        ends = np.array(ends)
        best, valid = (-1, float("-inf")), None
        for left in range(8):
            for right in range(left + 1, 9):
                edge = np.polyfit(middle_x[[left, right]], ends[[left, right]], 1)
                residual = np.abs(np.polyval(edge, middle_x) - ends)
                support = residual < height * 0.015
                score = (int(support.sum()), -float(residual[support].sum()))
                if score > best:
                    best, valid = score, support
        if best[0] < 5:
            raise ValueError("Table outline is unclear; use a photo showing the full table")
        edge = np.polyfit(middle_x[valid], ends[valid], 1)
        edge[1] += padding
        return edge

    top_edge = horizontal_edge(tops, -6)
    bottom_edge = horizontal_edge(bottoms, 6)

    def intersection(rule, edge):
        y = np.clip((edge[0] * rule[1] + edge[1]) / (1 - edge[0] * rule[0]), 0, height - 1)
        return float(np.polyval(rule, y) / scale), float(y / scale)

    corners = [intersection(rule, edge) for rule, edge in
               ((rules[0], top_edge), (rules[-1], top_edge),
                (rules[-1], bottom_edge), (rules[0], bottom_edge))]
    table = rectify(image, corners)
    # Map each detected divider through the same perspective transform.
    target = np.float32([[0, 0], [table.width - 1, 0],
                         [table.width - 1, table.height - 1], [0, table.height - 1]])
    matrix = cv2.getPerspectiveTransform(np.float32(corners), target)
    points = np.float32([[[np.polyval(rule, height / 2) / scale,
                          height / 2 / scale]] for rule in rules[1:-1]])
    columns = [0, *[round(float(point[0])) for point in cv2.perspectiveTransform(points, matrix)[:, 0]], table.width]
    validate_edges(columns, table.width, 9)
    return table, {"image_size": list(image.size), "corners": corners, "columns": columns}


def detect_rows(column_results, height):
    """Use transaction numbers and dates as row starts, preserving wrapped content."""
    anchors = []
    for column in range(3):
        for token in column_results[column]:
            text = token["text"].strip()
            if (column == 0 and text.isdigit()) or (column in (1, 2) and re.match(r"\d{1,2}/\d{1,2}/\d{4}", text)):
                anchors.append(token)
    if not anchors:
        raise ValueError("No transaction numbers or dates detected; cannot identify statement rows")
    # A date/time box can contain two lines; use single-line anchor height.
    line_height = min(statistics.median(t["box"][3] - t["box"][1] for t in anchors), height / 40)
    groups = []
    for token in sorted(anchors, key=lambda t: t["box"][1]):
        y = token["box"][1]
        if not groups or y - statistics.median(groups[-1]) > max(3, line_height * 0.75):
            groups.append([])
        groups[-1].append(y)
    # ponytail: rows need a readable number or date; unreadable anchors require a clearer photo.
    starts = [max(0, round(statistics.median(group) - line_height * 0.35)) for group in groups]
    rows = [*starts, height]
    if any(b - a < 3 for a, b in zip(rows, rows[1:])):
        raise ValueError("Detected transaction rows overlap")
    return rows


def create_ocr():
    os.environ.setdefault("PADDLE_PDX_CACHE_HOME", str(SCRIPT_DIR.parent / ".venv-ocr" / "model_cache"))
    from paddleocr import PaddleOCR, TextRecognition

    ocr = PaddleOCR(device="cpu", enable_mkldnn=False, cpu_threads=4,
                    text_detection_model_name="PP-OCRv6_medium_det",
                    text_recognition_model_name="PP-OCRv6_medium_rec",
                    use_doc_orientation_classify=False, use_doc_unwarping=False,
                    use_textline_orientation=False, text_rec_score_thresh=0.0)
    numbers = TextRecognition(model_name="PP-OCRv6_medium_rec", device="cpu",
                              enable_mkldnn=False, cpu_threads=4)
    return ocr, numbers


def read_tokens(image, ocr, scale=1):
    import cv2
    import numpy as np

    crop = np.asarray(image)[:, :, ::-1].copy()
    if scale > 1:
        crop = cv2.resize(crop, None, fx=scale, fy=scale, interpolation=cv2.INTER_CUBIC)
    crop = cv2.copyMakeBorder(crop, 12, 12, 12, 12, cv2.BORDER_CONSTANT, value=(255, 255, 255))
    tokens = []
    for prediction in ocr.predict(crop):
        for text, score, box in zip(prediction["rec_texts"], prediction["rec_scores"], prediction["rec_boxes"], strict=True):
            if text.strip():
                x0, y0, x1, y1 = (int(v) for v in box)
                tokens.append({"text": text, "score": float(score),
                               "box": [round((x0 - 12) / scale + 2), round((y0 - 12) / scale),
                                       round((x1 - 12) / scale + 2), round((y1 - 12) / scale)]})
    return tokens


def read_index_tokens(table, layout, recognizer):
    import cv2
    import numpy as np

    left, right = layout["columns"][:2]
    crops, boxes = [], []
    for top, bottom in zip(layout["rows"], layout["rows"][1:]):
        crop = np.asarray(table.crop((left + 3, top, right - 3, bottom)))[:, :, ::-1].copy()
        gray = cv2.cvtColor(crop, cv2.COLOR_BGR2GRAY)
        ink = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                    cv2.THRESH_BINARY_INV, 31, 10)
        _, _, stats, _ = cv2.connectedComponentsWithStats(ink)
        # Discard specks and tall divider fragments, then trim whitespace around the digits.
        letters = [(x, y, x + w, y + h) for x, y, w, h, area in stats[1:]
                   if area >= 5 and w >= 2 and h >= 4 and h < crop.shape[0] * 0.8 and h <= crop.shape[1]]
        if not letters:
            continue
        x0 = max(0, min(b[0] for b in letters) - 1)
        y0 = max(0, min(b[1] for b in letters) - 1)
        x1 = min(crop.shape[1], max(b[2] for b in letters) + 1)
        y1 = min(crop.shape[0], max(b[3] for b in letters) + 1)
        crops.append(crop[y0:y1, x0:x1])
        boxes.append([x0 + 3, y0 + top, x1 + 3, y1 + top])
    tokens = []
    if crops:
        for box, prediction in zip(boxes, recognizer.predict(crops, batch_size=8), strict=True):
            text = prediction["rec_text"].strip()
            if text:
                tokens.append({"text": text, "score": float(prediction["rec_score"]), "box": box})
    return tokens


def extract(table, layout, threshold, ocr, number_recognizer):
    column_results = [[]]
    # Read dates first; a full-height narrow index crop can merge many row numbers.
    for column in range(1, len(COLUMNS)):
        left, right = layout["columns"][column:column + 2]
        print(f"Reading column {column + 1}/8: {COLUMNS[column]}...", flush=True)
        image = table.crop((left + 2, 0, right - 2, table.height))
        column_results.append(read_tokens(image, ocr, scale=3 if image.width < 60 else 1))
    layout["rows"] = detect_rows(column_results, table.height)
    print(f"Reading {len(layout['rows']) - 1} row-number cells...", flush=True)
    column_results[0] = read_index_tokens(table, layout, number_recognizer)
    return assemble(column_results, layout["rows"], threshold)


def prepare_table(path):
    from PIL import Image, ImageOps

    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
    return detect_table(image)


def main():
    if not 0 <= MIN_CONFIDENCE <= 1:
        raise ValueError("MIN_CONFIDENCE must be between 0 and 1")
    filename = OUTPUT_FILENAME or "statement_results_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ".xlsx"
    output_name = Path(filename)
    if output_name.name != filename or output_name.suffix.lower() != ".xlsx":
        raise ValueError("OUTPUT_FILENAME must be an .xlsx filename without a folder path")
    output = SCRIPT_DIR / output_name
    if output.exists():
        raise FileExistsError(f"Output exists: {output}; change OUTPUT_FILENAME or leave it blank")
    selected_folder = Path(IMAGE_FOLDER)
    folder = selected_folder if selected_folder.is_absolute() else SCRIPT_DIR / selected_folder
    if not folder.exists() and selected_folder == Path("images"):
        folder.mkdir()
        print(f"Created {folder}. Put images there, then run this script again.")
        return
    if not folder.is_dir():
        raise ValueError(f"Image folder does not exist: {folder}; edit IMAGE_FOLDER")
    images = discover_images(folder)
    if not images:
        print(f"No supported images in {folder}")
        return
    if not (3, 10) <= sys.version_info[:2] <= (3, 13):
        raise RuntimeError(f"Python {sys.version.split()[0]} cannot run this PaddleOCR setup. Use 64-bit Python 3.12 in a separate virtual environment.")
    results, failures = [], []
    print("Loading PP-OCRv6 medium (CPU)...", flush=True)
    ocr, number_recognizer = create_ocr()
    for number, path in enumerate(images, 1):
        print(f"[{number}/{len(images)}] {path.name}", flush=True)
        try:
            table, current = prepare_table(path)
            result = extract(table, current, MIN_CONFIDENCE, ocr, number_recognizer)
            result["source"] = path.name
            results.append(result)
        except KeyboardInterrupt:
            failures.append({"source_file": path.name, "error": "Stopped by user"})
            failures.extend({"source_file": remaining.name, "error": "Not processed: batch interrupted"}
                            for remaining in images[number:])
            print("Interrupted; saving completed images.")
            break
        except Exception as error:
            print(f"Failed {path.name}: {error}")
            failures.append({"source_file": path.name, "error": "Failed: " + str(error)})
    write_excel(output, results, failures)
    flagged = sum(len({issue["row"] for issue in result["issues"] if "row" in issue}) for result in results)
    print(f"Saved {sum(len(result['rows']) for result in results)} rows from {len(results)} images to {output}")
    print(f"{flagged} transaction rows flagged; {len(failures)} images failed or were not processed. Filter the flagged column to review.")


if __name__ == "__main__":
    interpreter = SCRIPT_DIR.parent / ".venv-ocr" / "Scripts" / "python.exe"
    if interpreter.is_file() and Path(sys.executable).resolve() != interpreter.resolve():
        import subprocess

        raise SystemExit(subprocess.call([str(interpreter), str(Path(__file__).resolve())]))
    main()
