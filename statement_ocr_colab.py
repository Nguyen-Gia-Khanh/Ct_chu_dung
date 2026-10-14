"""Statement OCR Pipeline for Google Colab (PaddleOCR-VL-1.6 & PP-OCRv6).

Extracts transaction records from bank statement photos in Google Drive,
runs validation checks (missing amounts, invalid index, sequence gaps),
and exports the records to Google Sheets with cell-level issue highlighting
and an optional backup Excel file.
"""

import bisect
from datetime import datetime
from html import unescape
import os
from pathlib import Path
import re
import statistics
import sys
from uuid import uuid4
import cv2
import numpy as np
from PIL import Image, ImageOps

COLUMNS = ("index", "time", "date", "id", "content", "money_out", "money_in", "total_raw")
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}

ROW_ISSUE_COLUMNS = {
    "no_amount_detected": ("money_out", "money_in"),
    "both_amount_columns_populated": ("money_out", "money_in"),
    "index_gap_or_duplicate": ("index",),
}


def discover_images(folder):
    """Find and naturally sort image files."""
    def natural_key(path):
        return tuple((1, int(part)) if part.isdigit() else (0, part.casefold())
                     for part in re.split(r"(\d+)", path.name))
    return sorted((p for p in folder.iterdir()
                   if p.is_file() and not p.name.startswith(".") and p.suffix.lower() in IMAGE_SUFFIXES),
                  key=natural_key)


def parse_html_table(html_str):
    """Extract rows and cells from an HTML table string."""
    rows = []
    tr_matches = re.findall(r"<tr[^>]*>(.*?)</tr>", html_str, flags=re.DOTALL | re.IGNORECASE)
    for tr in tr_matches:
        cells = re.findall(r"<(?:td|th)[^>]*>(.*?)</(?:td|th)>", tr, flags=re.DOTALL | re.IGNORECASE)
        clean_cells = [unescape(re.sub(r"<[^>]+>", "", c)).strip() for c in cells]
        if any(clean_cells):
            rows.append(clean_cells)
    return rows


def validate_and_flag_records(records, min_confidence=0.80, default_score=0.95):
    """Validate records, detect issues, and record columns requiring review."""
    issues = []
    review_columns = {}
    confidence = []

    for row_number, record in enumerate(records, 1):
        scores = {col: default_score if record.get(col) else None for col in COLUMNS}
        confidence.append(scores)

        # 1. Numeric index check
        if not record["index"].isdigit():
            issues.append({"row": row_number, "column": "index", "issue": "missing_or_non_numeric_index"})
            review_columns.setdefault(row_number, set()).add("index")

        # 2. Text fields check
        for name in ("time", "date", "id", "content", "total_raw"):
            if not record[name]:
                issues.append({"row": row_number, "column": name, "issue": "no_text_detected"})
                review_columns.setdefault(row_number, set()).add(name)

        # 3. Amounts check
        if not record["money_out"] and not record["money_in"]:
            issues.append({"row": row_number, "issue": "no_amount_detected"})
            review_columns.setdefault(row_number, set()).update(ROW_ISSUE_COLUMNS["no_amount_detected"])
        if record["money_out"] and record["money_in"]:
            issues.append({"row": row_number, "issue": "both_amount_columns_populated"})
            review_columns.setdefault(row_number, set()).update(ROW_ISSUE_COLUMNS["both_amount_columns_populated"])

    # 4. Sequence continuity check
    indices = [int(r["index"]) if r["index"].isdigit() else None for r in records]
    for row, (previous, current) in enumerate(zip(indices, indices[1:]), 2):
        if previous is not None and current is not None and current != previous + 1:
            issues.append({"row": row, "issue": "index_gap_or_duplicate", "previous": previous, "current": current})
            review_columns.setdefault(row, set()).add("index")

    return {"rows": records, "confidence": confidence, "issues": issues, "review_columns": review_columns}


# --- Engine 1: PaddleOCR-VL-1.6 (Vision-Language) ---
def create_vl_pipeline():
    """Create SOTA PaddleOCR-VL-1.6 vision-language pipeline."""
    # Fix PaddleX duplicate initialization bug in Colab/Jupyter
    try:
        import paddlex.repo_manager.core as pm_core
        if not hasattr(pm_core, "_safe_pdx_patched"):
            _orig_pm_init = pm_core.initialize
            pm_core.initialize = lambda *a, **kw: None if pm_core._GlobalContext.is_initialized() else _orig_pm_init(*a, **kw)
            pm_core._safe_pdx_patched = True
    except Exception:
        pass

    from paddlex import create_pipeline
    import paddle
    if paddle.is_compiled_with_cuda():
        paddle.set_device("gpu:0")
        device = "gpu:0"
    else:
        device = "cpu"
    print(f"Initializing PaddleOCR-VL-1.6 Pipeline on {device.upper()}...")
    return create_pipeline(pipeline="PaddleOCR-VL-1.6", device=device)


def extract_with_vl_16(image_path, pipeline, min_confidence=0.80):
    """Run PaddleOCR-VL-1.6 to parse tables directly without rigid line rules."""
    output = pipeline.predict(input=str(image_path))
    extracted_rows = []

    for res in output:
        # Check table_res_list
        table_list = res.get("table_res_list", [])
        if table_list:
            for table_res in table_list:
                html_code = ""
                if hasattr(table_res, "html") and isinstance(table_res.html, dict):
                    html_code = table_res.html.get("pred", "")
                elif isinstance(table_res, dict):
                    html_code = table_res.get("pred_html", "") or table_res.get("html", "")
                if html_code:
                    extracted_rows.extend(parse_html_table(html_code))

        # Check parsing_res_list blocks
        if not extracted_rows:
            blocks = res.get("parsing_res_list", [])
            for block in blocks:
                label = getattr(block, "label", None) or (block.get("label") if isinstance(block, dict) else "")
                content = getattr(block, "content", "") or (block.get("content") if isinstance(block, dict) else "")
                if label == "table" and content:
                    extracted_rows.extend(parse_html_table(content))

    if not extracted_rows:
        raise ValueError("PaddleOCR-VL-1.6 did not detect any table in this photo.")

    # Filter out header row if first row has non-digit index (e.g. 'STT')
    if extracted_rows and not extracted_rows[0][0].isdigit():
        extracted_rows = extracted_rows[1:]

    records = []
    for raw_row in extracted_rows:
        cells = (raw_row + [""] * 8)[:8]
        record = {name: cell.strip() for name, cell in zip(COLUMNS, cells)}
        records.append(record)

    return validate_and_flag_records(records, min_confidence)


# --- Engine 2: PP-OCRv6 Geometric ---
def create_v6_ocr():
    from paddleocr import PaddleOCR, TextRecognition
    import paddle
    device = "gpu" if paddle.is_compiled_with_cuda() else "cpu"
    print(f"Initializing PP-OCRv6 on {device.upper()}...")
    ocr_kwargs = {
        "device": device,
        "text_detection_model_name": "PP-OCRv6_medium_det",
        "text_recognition_model_name": "PP-OCRv6_medium_rec",
        "use_doc_orientation_classify": False,
        "use_doc_unwarping": False,
        "use_textline_orientation": False,
        "text_rec_score_thresh": 0.0,
    }
    rec_kwargs = {"model_name": "PP-OCRv6_medium_rec", "device": device}
    if device == "cpu":
        ocr_kwargs["enable_mkldnn"] = False
        ocr_kwargs["cpu_threads"] = 4
        rec_kwargs["enable_mkldnn"] = False
        rec_kwargs["cpu_threads"] = 4
    return PaddleOCR(**ocr_kwargs), TextRecognition(**rec_kwargs)


def rectify(image, corners):
    points = np.asarray(corners, dtype=np.float32)
    if points.shape != (4, 2) or not cv2.isContourConvex(points.astype(np.int32)):
        raise ValueError("Could not determine a convex table outline")
    width = round(max(np.linalg.norm(points[1] - points[0]), np.linalg.norm(points[2] - points[3])))
    height = round(max(np.linalg.norm(points[3] - points[0]), np.linalg.norm(points[2] - points[1])))
    if min(width, height) < 30:
        raise ValueError("Detected table is too small")
    target = np.float32([[0, 0], [width - 1, 0], [width - 1, height - 1], [0, height - 1]])
    matrix = cv2.getPerspectiveTransform(points, target)
    return Image.fromarray(cv2.warpPerspective(np.asarray(image), matrix, (width, height)))


def detect_table(image):
    scale = min(1, 1200 / max(image.size))
    photo = cv2.resize(np.asarray(image), (round(image.width * scale), round(image.height * scale)))
    height, width = photo.shape[:2]
    gray = cv2.cvtColor(photo, cv2.COLOR_RGB2GRAY)
    binary = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                   cv2.THRESH_BINARY_INV, 31, 10)
    binary[cv2.cvtColor(photo, cv2.COLOR_RGB2HSV)[:, :, 1] > 80] = 0
    binary = cv2.morphologyEx(binary, cv2.MORPH_OPEN,
                             np.ones((max(5, height // 80), 1), np.uint8))
    lines = cv2.HoughLinesP(binary, 1, np.pi / 720, threshold=max(20, height // 27),
                           minLineLength=height // 10, maxLineGap=height // 30)
    candidates = []
    for x0, y0, x1, y1 in (() if lines is None else lines.reshape(-1, 4)):
        if abs(y1 - y0) > height * 0.15 and abs(x1 - x0) < abs(y1 - y0) * 0.25:
            middle = x0 + (height / 2 - y0) * (x1 - x0) / (y1 - y0)
            candidates.append((middle, (x0, y0, x1, y1)))
    groups = []
    for middle, line in sorted(candidates):
        if not groups or middle - statistics.mean(item[0] for item in groups[-1]) > max(3, width * 0.008):
            groups.append([])
        groups[-1].append((middle, line))
    if len(groups) != 9:
        raise ValueError(f"Found {len(groups)} column boundaries; expected 9 for this eight-column statement.")
    rules, tops, bottoms = [], [], []
    for group in groups:
        points = np.array([(x, y) for _, (x0, y0, x1, y1) in group
                           for x, y in ((x0, y0), (x1, y1))], dtype=float)
        rules.append(np.polyfit(points[:, 1], points[:, 0], 1))
        tops.append(float(points[:, 1].min()))
        bottoms.append(float(points[:, 1].max()))
    middle_x = np.array([np.polyval(rule, height / 2) for rule in rules])

    def horizontal_edge(ends, padding):
        ends = np.array(ends)
        best, valid = (-1, float("-inf")), None
        for left in range(8):
            for right in range(left + 1, 9):
                edge = np.polyfit(middle_x[[left, right]], ends[[left, right]], 1)
                residual = np.abs(np.polyval(edge, middle_x) - ends)
                support = residual < height * 0.015
                score = (int(support.sum()), -float(residual[support].sum()))
                if score > best:
                    best, valid = score, support
        if best[0] < 5:
            raise ValueError("Table outline is unclear; use a photo showing the full table")
        edge = np.polyfit(middle_x[valid], ends[valid], 1)
        edge[1] += padding
        return edge

    top_edge = horizontal_edge(tops, -6)
    bottom_edge = horizontal_edge(bottoms, 6)

    def intersection(rule, edge):
        y = np.clip((edge[0] * rule[1] + edge[1]) / (1 - edge[0] * rule[0]), 0, height - 1)
        return float(np.polyval(rule, y) / scale), float(y / scale)

    corners = [intersection(rule, edge) for rule, edge in
               ((rules[0], top_edge), (rules[-1], top_edge),
                (rules[-1], bottom_edge), (rules[0], bottom_edge))]
    table = rectify(image, corners)
    target = np.float32([[0, 0], [table.width - 1, 0],
                         [table.width - 1, table.height - 1], [0, table.height - 1]])
    matrix = cv2.getPerspectiveTransform(np.float32(corners), target)
    points = np.float32([[[np.polyval(rule, height / 2) / scale,
                          height / 2 / scale]] for rule in rules[1:-1]])
    columns = [0, *[round(float(point[0])) for point in cv2.perspectiveTransform(points, matrix)[:, 0]], table.width]
    return table, {"image_size": list(image.size), "corners": corners, "columns": columns}


def detect_rows(column_results, height):
    anchors = []
    for column in range(3):
        for token in column_results[column]:
            text = token["text"].strip()
            if (column == 0 and text.isdigit()) or (column in (1, 2) and re.match(r"\d{1,2}/\d{1,2}/\d{4}", text)):
                anchors.append(token)
    if not anchors:
        raise ValueError("No transaction numbers or dates detected; cannot identify statement rows")
    line_height = min(statistics.median(t["box"][3] - t["box"][1] for t in anchors), height / 40)
    groups = []
    for token in sorted(anchors, key=lambda t: t["box"][1]):
        y = token["box"][1]
        if not groups or y - statistics.median(groups[-1]) > max(3, line_height * 0.75):
            groups.append([])
        groups[-1].append(y)
    starts = [max(0, round(statistics.median(group) - line_height * 0.35)) for group in groups]
    rows = [*starts, height]
    return rows


def read_tokens(image, ocr, scale=1):
    crop = np.asarray(image)[:, :, ::-1].copy()
    if scale > 1:
        crop = cv2.resize(crop, None, fx=scale, fy=scale, interpolation=cv2.INTER_CUBIC)
    crop = cv2.copyMakeBorder(crop, 12, 12, 12, 12, cv2.BORDER_CONSTANT, value=(255, 255, 255))
    tokens = []
    for prediction in ocr.predict(crop):
        for text, score, box in zip(prediction["rec_texts"], prediction["rec_scores"], prediction["rec_boxes"], strict=True):
            if text.strip():
                x0, y0, x1, y1 = (int(v) for v in box)
                tokens.append({"text": text, "score": float(score),
                               "box": [round((x0 - 12) / scale + 2), round((y0 - 12) / scale),
                                       round((x1 - 12) / scale + 2), round((y1 - 12) / scale)]})
    return tokens


def read_index_tokens(table, layout, recognizer):
    left, right = layout["columns"][:2]
    crops, boxes = [], []
    for top, bottom in zip(layout["rows"], layout["rows"][1:]):
        crop = np.asarray(table.crop((left + 3, top, right - 3, bottom)))[:, :, ::-1].copy()
        gray = cv2.cvtColor(crop, cv2.COLOR_BGR2GRAY)
        ink = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                    cv2.THRESH_BINARY_INV, 31, 10)
        _, _, stats, _ = cv2.connectedComponentsWithStats(ink)
        letters = [(x, y, x + w, y + h) for x, y, w, h, area in stats[1:]
                   if area >= 5 and w >= 2 and h >= 4 and h < crop.shape[0] * 0.8 and h <= crop.shape[1]]
        if not letters:
            continue
        x0 = max(0, min(b[0] for b in letters) - 1)
        y0 = max(0, min(b[1] for b in letters) - 1)
        x1 = min(crop.shape[1], max(b[2] for b in letters) + 1)
        y1 = min(crop.shape[0], max(b[3] for b in letters) + 1)
        crops.append(crop[y0:y1, x0:x1])
        boxes.append([x0 + 3, y0 + top, x1 + 3, y1 + top])
    tokens = []
    if crops:
        for box, prediction in zip(boxes, recognizer.predict(crops, batch_size=8), strict=True):
            text = prediction["rec_text"].strip()
            if text:
                tokens.append({"text": text, "score": float(prediction["rec_score"]), "box": box})
    return tokens


def extract_with_v6(table, layout, threshold, ocr, number_recognizer):
    column_results = [[]]
    for column in range(1, len(COLUMNS)):
        left, right = layout["columns"][column:column + 2]
        image = table.crop((left + 2, 0, right - 2, table.height))
        column_results.append(read_tokens(image, ocr, scale=3 if image.width < 60 else 1))
    layout["rows"] = detect_rows(column_results, table.height)
    column_results[0] = read_index_tokens(table, layout, number_recognizer)

    cells = [[[] for _ in COLUMNS] for _ in range(len(layout["rows"]) - 1)]
    for column, tokens in enumerate(column_results):
        for token in tokens:
            x0, y0, x1, y1 = token["box"]
            row = bisect.bisect_right(layout["rows"], (y0 + y1) / 2) - 1
            if 0 <= row < len(cells):
                cells[row][column].append(token)

    records = []
    for row in cells:
        record = {}
        for name, tokens in zip(COLUMNS, row):
            tokens.sort(key=lambda t: (round(t["box"][1] / 5), t["box"][0]))
            record[name] = " ".join(t["text"] for t in tokens).strip()
        records.append(record)

    return validate_and_flag_records(records, threshold)


# --- Google Sheets & Excel Exporters ---
def export_to_google_sheets(gc, results, failures, sheet_title):
    headers = list(COLUMNS) + ["flagged"]
    all_rows = [headers]
    review_cells = []

    for result in results:
        flags = {}
        for issue in result.get("issues", []):
            message = issue["issue"].replace("_", " ")
            if issue.get("column"):
                message = issue["column"] + ": " + message
            if issue["issue"] == "low_confidence":
                message += f" {issue['score']:.2f} < {issue['threshold']:.2f}"
            if "row" not in issue:
                message += ": " + issue.get("token", {}).get("text", "")
            flags.setdefault(issue.get("row"), []).append(message)

        review_columns = result.get("review_columns", {})

        for number, record in enumerate(result.get("rows", []), 1):
            current_row_idx = len(all_rows)
            messages = flags.get(number, [])
            flagged = result["source"] + ": " + "; ".join(dict.fromkeys(messages)) if messages else ""
            row_data = [str(record.get(col) or "") for col in COLUMNS] + [flagged]
            all_rows.append(row_data)

            # Record affected cells for red/pink (#F4CCCC) highlight
            for affected_col in review_columns.get(number, ()):
                if affected_col in COLUMNS:
                    col_idx = COLUMNS.index(affected_col)
                    review_cells.append((current_row_idx, col_idx))

        if flags.get(None):
            all_rows.append([""] * len(COLUMNS) + [result["source"] + ": " + "; ".join(dict.fromkeys(flags[None]))])

    for failure in failures:
        all_rows.append([""] * len(COLUMNS) + [failure["source_file"] + ": " + failure["error"]])

    print(f"Creating Google Sheet: '{sheet_title}'...")
    sh = gc.create(sheet_title)
    ws = sh.sheet1
    try:
        ws.update_title("Statement Results")
    except Exception:
        pass

    num_rows = max(len(all_rows) + 50, 100)
    num_cols = max(len(headers), 10)
    try:
        ws.resize(rows=num_rows, cols=num_cols)
    except Exception:
        pass

    try:
        ws.update(range_name='A1', values=all_rows)
    except TypeError:
        ws.update('A1', all_rows)

    col_widths = [80, 180, 130, 180, 500, 150, 150, 150, 600]
    requests = []

    # 1. Custom column widths
    for i, width in enumerate(col_widths):
        requests.append({
            "updateDimensionProperties": {
                "range": {
                    "sheetId": ws.id,
                    "dimension": "COLUMNS",
                    "startIndex": i,
                    "endIndex": i + 1,
                },
                "properties": {"pixelSize": width},
                "fields": "pixelSize",
            }
        })

    # 2. Header styling: dark slate blue (#243746), bold white text, centered
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": ws.id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": len(headers),
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": {"red": 36 / 255, "green": 55 / 255, "blue": 70 / 255},
                    "textFormat": {
                        "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0},
                        "bold": True,
                    },
                    "horizontalAlignment": "CENTER",
                    "verticalAlignment": "MIDDLE",
                }
            },
            "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment)",
        }
    })

    # 3. Freeze row 1
    requests.append({
        "updateSheetProperties": {
            "properties": {
                "sheetId": ws.id,
                "gridProperties": {
                    "frozenRowCount": 1,
                },
            },
            "fields": "gridProperties.frozenRowCount",
        }
    })

    # 4. Highlight flagged rows with soft yellow (#FFF2CC)
    requests.append({
        "addConditionalFormatRule": {
            "rule": {
                "ranges": [{
                    "sheetId": ws.id,
                    "startRowIndex": 1,
                    "endRowIndex": len(all_rows),
                    "startColumnIndex": 0,
                    "endColumnIndex": len(headers),
                }],
                "booleanRule": {
                    "condition": {
                        "type": "CUSTOM_FORMULA",
                        "values": [{"userEnteredValue": '=$I2<>""'}],
                    },
                    "format": {
                        "backgroundColor": {"red": 255 / 255, "green": 242 / 255, "blue": 204 / 255},
                    },
                },
            },
            "index": 0,
        }
    })

    # 5. Cell-level highlight: soft pink/red (#F4CCCC) for specific problem cells
    for r_idx, c_idx in review_cells:
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": ws.id,
                    "startRowIndex": r_idx,
                    "endRowIndex": r_idx + 1,
                    "startColumnIndex": c_idx,
                    "endColumnIndex": c_idx + 1,
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": {"red": 244 / 255, "green": 204 / 255, "blue": 204 / 255},
                    }
                },
                "fields": "userEnteredFormat.backgroundColor",
            }
        })

    try:
        sh.batch_update({"requests": requests})
    except Exception as e:
        print(f"Formatting note: {e}")

    return sh, len(all_rows) - 1


def export_to_excel(path, results, failures=()):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append((*COLUMNS, "flagged"))
    review_fill = PatternFill("solid", fgColor="F4CCCC")

    for result in results:
        flags = {}
        for issue in result.get("issues", []):
            message = issue["issue"].replace("_", " ")
            if issue.get("column"):
                message = issue["column"] + ": " + message
            if issue["issue"] == "low_confidence":
                message += f" {issue['score']:.2f} < {issue['threshold']:.2f}"
            if "row" not in issue:
                message += ": " + issue.get("token", {}).get("text", "")
            flags.setdefault(issue.get("row"), []).append(message)

        review_columns = result.get("review_columns", {})

        for number, record in enumerate(result.get("rows", []), 1):
            messages = flags.get(number, [])
            flagged = result["source"] + ": " + "; ".join(dict.fromkeys(messages)) if messages else None
            data.append((*[record.get(name) or None for name in COLUMNS], flagged))

            for name in review_columns.get(number, ()):
                if name in COLUMNS:
                    data.cell(data.max_row, COLUMNS.index(name) + 1).fill = review_fill

            data.row_dimensions[data.max_row].height = min(160, max(30, 15 * (len(record.get("content", "")) // 75 + 1)))

        if flags.get(None):
            data.append((*[None for _ in COLUMNS], result["source"] + ": " + "; ".join(dict.fromkeys(flags[None]))))

    for failure in failures:
        data.append((*[None for _ in COLUMNS], failure["source_file"] + ": " + failure["error"]))

    data.freeze_panes = "A2"
    data.auto_filter.ref = data.dimensions
    data.row_dimensions[1].height = 24
    for cell in data[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for column, width in enumerate((9, 24, 16, 24, 78, 20, 20, 20, 100), 1):
        data.column_dimensions[data.cell(1, column).column_letter].width = width
    for row in data.iter_rows(min_row=2):
        for cell in row:
            cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row[-1].value:
            row[-1].fill = PatternFill("solid", fgColor="FFF2CC")
    workbook.save(path)
